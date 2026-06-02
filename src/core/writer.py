from pathlib import Path
from loguru import logger
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.styles import Font, Alignment, Border, Side
from .models import SpecDocument

# Настройки вывода
WRITER_CONFIG = {
    "sheet_name": "Спецификация",
    "headers": ['№ п/п', 'Наименование', 'Ед. изм.', 'Кол-во', 'Сумма, руб'],
    "column_widths": {'A': 12, 'B': 80, 'C': 12, 'D': 10, 'E': 18},
    "total_text": "ИТОГО по разделу",
    "grand_total_text": "ОБЩИЙ ИТОГ",
    "enable_wrap_text": True,
    "enable_auto_height": True,
}


def create_result_sheet(wb: Workbook, sheet_name: str = None) -> Worksheet:
    """Создаёт новый лист для результата в чистом workbook"""
    if sheet_name is None:
        sheet_name = WRITER_CONFIG["sheet_name"]
    return wb.create_sheet(sheet_name)


def write_spec_to_sheet(wb: Workbook, doc: SpecDocument) -> None:
    """Записывает спецификацию на новый лист в чистом workbook"""
    ws = create_result_sheet(wb)

    # Стили
    title_font = Font(bold=True, size=14, name='Arial')
    header_font = Font(bold=True, size=11, name='Arial')
    section_font = Font(bold=True, size=11, name='Arial')
    total_font = Font(bold=True, size=11, name='Arial')
    normal_font = Font(size=10, name='Arial')

    center_align = Alignment(horizontal='center', vertical='center')
    left_top_align = Alignment(horizontal='left', vertical='top', wrap_text=WRITER_CONFIG["enable_wrap_text"])
    right_align = Alignment(horizontal='right', vertical='center')

    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    row = 1

    # Заголовок
    ws.merge_cells(f'A{row}:E{row}')
    ws.cell(row, 1).value = "Техническая спецификация"
    ws.cell(row, 1).font = title_font
    ws.cell(row, 1).alignment = center_align
    row += 2

    # Название изделия
    if doc.header and doc.header.equipment_type:
        ws.merge_cells(f'A{row}:E{row}')
        ws.cell(row, 1).value = doc.header.equipment_type
        ws.cell(row, 1).font = Font(bold=True, size=12)
        ws.cell(row, 1).alignment = center_align
        row += 2

    # Шапка таблицы
    for col, header in enumerate(WRITER_CONFIG["headers"], 1):
        cell = ws.cell(row, col, header)
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
    row += 1

    # Заполнение данных
    for section in doc.sections:
        # Заголовок секции
        ws.cell(row, 1).value = str(section.number)
        ws.cell(row, 2).value = section.title
        ws.cell(row, 1).font = section_font
        ws.cell(row, 2).font = section_font
        ws.cell(row, 1).alignment = center_align
        ws.cell(row, 2).alignment = center_align
        for col in range(1, 6):
            ws.cell(row, col).border = thin_border
        row += 1

        # Позиции
        for item in section.items:
            # Номер
            cell_num = ws.cell(row, 1, item.number)
            cell_num.border = thin_border
            cell_num.alignment = center_align
            cell_num.font = normal_font

            # Наименование
            cell_name = ws.cell(row, 2, item.name)
            cell_name.border = thin_border
            cell_name.font = normal_font
            cell_name.alignment = left_top_align

            # Автовысота
            if WRITER_CONFIG["enable_auto_height"] and item.name and ('\n' in item.name):
                lines = item.name.count('\n') + 1
                ws.row_dimensions[row].height = 12.75 * lines

            # Ед. изм
            cell_unit = ws.cell(row, 3, item.unit)
            cell_unit.border = thin_border
            cell_unit.alignment = center_align
            cell_unit.font = normal_font

            # Количество
            cell_qty = ws.cell(row, 4, item.quantity)
            cell_qty.border = thin_border
            cell_qty.alignment = center_align
            cell_qty.font = normal_font

            # Сумма
            cell_total = ws.cell(row, 5, item.total)
            cell_total.border = thin_border
            cell_total.alignment = right_align
            cell_total.font = normal_font
            cell_total.number_format = '#,##0.00'

            row += 1

        # Итог по разделу
        if section.section_total > 0:
            ws.merge_cells(f'A{row}:D{row}')
            cell = ws.cell(row, 1, WRITER_CONFIG["total_text"])
            cell.border = thin_border
            cell.font = total_font
            cell.alignment = right_align

            cell_total = ws.cell(row, 5, section.section_total)
            cell_total.border = thin_border
            cell_total.font = total_font
            cell_total.alignment = right_align
            cell_total.number_format = '#,##0.00'
            row += 1

        row += 1

    # Общий итог
    if doc.grand_total > 0:
        ws.merge_cells(f'A{row}:D{row}')
        cell = ws.cell(row, 1, WRITER_CONFIG["grand_total_text"])
        cell.border = thin_border
        cell.font = Font(bold=True, size=12)
        cell.alignment = right_align

        cell_total = ws.cell(row, 5, doc.grand_total)
        cell_total.border = thin_border
        cell_total.font = Font(bold=True, size=12)
        cell_total.alignment = right_align
        cell_total.number_format = '#,##0.00'

    # Установка ширины колонок
    for col_letter, width in WRITER_CONFIG["column_widths"].items():
        ws.column_dimensions[col_letter].width = width

    # Включаем перенос текста для колонки B
    for r in range(1, row + 1):
        cell = ws.cell(r, 2)
        if cell.value:
            cell.alignment = left_top_align

    logger.info(
        f"Создан лист '{WRITER_CONFIG['sheet_name']}' с {doc.total_items} позициями, итого: {doc.grand_total:.2f}")


def save_workbook(wb: Workbook, file_path: Path) -> None:
    """Сохраняет workbook"""
    try:
        file_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(file_path))
        logger.info(f"Файл сохранён: {file_path}")
    except Exception as e:
        logger.error(f"Ошибка сохранения: {e}")
        raise
