from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from loguru import logger

from .models import SpecDocument


def set_cell_border(cell, border_size=1):
    """
    Установка границ для ячейки таблицы.
    """

    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()

    for border_name in ['top', 'left', 'bottom', 'right']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), str(border_size))
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), '000000')
        tcPr.append(border)


def merge_cells(row, start_col, end_col):
    """
    Объединение ячеек в строке от start_col до end_col.
    """

    if start_col == end_col:
        return
    cell = row.cells[start_col]
    cell.merge(row.cells[end_col])


def set_cell_text(cell, text, bold=False, alignment=None):
    """
    Установка текста в ячейку с поддержкой переносов строк.
    """

    if not text:
        cell.text = ""
        return

    cell.text = ""
    lines = str(text).split('\n')

    for i, line in enumerate(lines):
        line = line.strip()
        if not line:
            continue
        if i == 0:
            run = cell.paragraphs[0].add_run(line)
        else:
            p = cell.add_paragraph()
            run = p.add_run(line)
        run.bold = bold
        if alignment:
            if i == 0:
                cell.paragraphs[0].alignment = alignment
            else:
                p.alignment = alignment


def create_word_specification(doc: SpecDocument, output_path: Path) -> None:
    """
    Создание Word-документа.
    """

    document = Document()

    style = document.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(10)

    title = document.add_heading('Техническая спецификация', level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if doc.header and doc.header.equipment_type:
        subtitle = document.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = subtitle.add_run(doc.header.equipment_type)
        run.bold = True
        run.font.size = Pt(12)

    document.add_paragraph()

    table = document.add_table(rows=1, cols=5)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    table.columns[0].width = Cm(2)
    table.columns[1].width = Cm(10)
    table.columns[2].width = Cm(2)
    table.columns[3].width = Cm(2)
    table.columns[4].width = Cm(3)

    headers = ['№ п/п', 'Наименование', 'Ед. изм.', 'Кол-во', 'Сумма, руб']
    header_row = table.rows[0]
    for i, header in enumerate(headers):
        cell = header_row.cells[i]
        set_cell_text(cell, header, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_border(cell)

    for section in doc.sections:
        row = table.add_row()
        set_cell_text(row.cells[0], str(section.number), bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[1], section.title, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        for col in [2, 3, 4]:
            row.cells[col].text = ""
        for cell in row.cells:
            set_cell_border(cell)

        for item in section.items:
            row = table.add_row()
            set_cell_text(row.cells[0], item.number, alignment=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_text(row.cells[1], item.name, alignment=WD_ALIGN_PARAGRAPH.LEFT)
            set_cell_text(row.cells[2], item.unit, alignment=WD_ALIGN_PARAGRAPH.CENTER)
            qty_str = str(int(item.quantity)) if item.quantity == int(item.quantity) else str(item.quantity)
            set_cell_text(row.cells[3], qty_str, alignment=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_text(row.cells[4], f"{item.total:,.2f} ₽", alignment=WD_ALIGN_PARAGRAPH.RIGHT)
            for cell in row.cells:
                set_cell_border(cell)

        if section.section_total > 0:
            row = table.add_row()
            merge_cells(row, 0, 3)
            set_cell_text(row.cells[0], "ИТОГО по разделу", bold=True, alignment=WD_ALIGN_PARAGRAPH.RIGHT)
            set_cell_text(row.cells[4], f"{section.section_total:,.2f} ₽", bold=True,
                          alignment=WD_ALIGN_PARAGRAPH.RIGHT)
            for cell in row.cells:
                set_cell_border(cell)

    if doc.grand_total > 0:
        row = table.add_row()
        merge_cells(row, 0, 3)
        set_cell_text(row.cells[0], "ОБЩИЙ ИТОГ", bold=True, alignment=WD_ALIGN_PARAGRAPH.RIGHT)
        set_cell_text(row.cells[4], f"{doc.grand_total:,.2f} ₽", bold=True, alignment=WD_ALIGN_PARAGRAPH.RIGHT)
        for cell in row.cells:
            set_cell_border(cell)

    document.save(str(output_path))
    logger.info(f"Word-документ сохранён: {output_path}")


def save_as_word(doc: SpecDocument, output_path: Path) -> None:
    """
    Сохранение Word-документа.
    """

    create_word_specification(doc, output_path)
