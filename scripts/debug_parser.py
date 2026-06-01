from openpyxl import load_workbook

file_path = "/home/amshegar/Загрузки/Attachments_hhz@gknzo.ru_2026-05-27_15-57-58/2055_ ЯКНО-ВВ(ВК)-6кВ_Соврудник.xlsm"
wb = load_workbook(file_path, data_only=True)
ws = wb.worksheets[0]

print(f"Лист: {ws.title}")
print("=" * 80)

for row_idx in range(25, 50):
    row = ws[row_idx]
    if row:
        a = row[0].value if row[0] else ""
        b = row[1].value if row[1] else ""
        c = row[2].value if row[2] else ""
        d = row[3].value if row[3] else ""
        e = row[4].value if row[4] else ""
        h = row[7].value if row[7] else ""
        print(f"Строка {row_idx}: A='{a}', B='{b}', C='{c}', D='{d}', E='{e}', H='{h}'")