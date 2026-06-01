#!/usr/bin/env python3
from pathlib import Path
from openpyxl import load_workbook

file_path = "/home/amshegar/Загрузки/Attachments_hhz@gknzo.ru_2026-05-27_15-57-58/2055_ ЯКНО-ВВ(ВК)-6кВ_Соврудник.xlsm"
wb = load_workbook(file_path, data_only=True)
ws = wb.worksheets[0]

# Найдем строку с Механическими блокировками
for row_idx in range(50, 70):
    row = ws[row_idx]
    if row and row[2] and row[2].value:
        val = str(row[2].value)
        if "Механические блокировки" in val:
            print(f"Строка {row_idx}:")
            print(f"  Значение: {repr(val)}")
            print(f"  Символы: {[ord(c) for c in val if ord(c) == 10]}")
            print(f"  Содержит \\n: {'\\n' in val}")
            print(f"  Содержит chr(10): {chr(10) in val}")
            break