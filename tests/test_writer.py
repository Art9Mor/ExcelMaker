from openpyxl import Workbook

from src.core.models import SpecDocument, Section, SpecItem
from src.core.writer import write_spec_to_sheet


def make_workbook_with_template() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Лист1"

    ws.append(["Структура", "Скрыть строку, символы /*", "№ п/п", "Наименование", "Цена", "Ед. изм.", "Кол-во", "Стоимость", "Стоимость с наценкой", "Ед. изм."])
    ws.append(["ИТОГ", "", "", "", "", "", "", "", "", ""])
    return wb


def make_document() -> SpecDocument:
    doc = SpecDocument(source_sheet="Лист1", target_sheet="Лист1")
    section = Section(number=1, title="Корпус")
    section.items.append(SpecItem(number="1.1", name="Каркас ЯКНО", unit="шт", quantity=1))
    section.items.append(SpecItem(number="1.2", name="Рама опорная", unit="шт", quantity=1))
    doc.sections.append(section)
    return doc


def test_write_spec_to_sheet_writes_sections_and_items():
    wb = make_workbook_with_template()
    doc = make_document()

    write_spec_to_sheet(wb, doc)

    ws = wb["Лист1"]

    assert ws.cell(2, 3).value == "1"
    assert ws.cell(2, 4).value == "Корпус"
    assert ws.cell(3, 3).value == "1.1"
    assert ws.cell(3, 4).value == "Каркас ЯКНО"
    assert ws.cell(3, 6).value == "шт"
    assert ws.cell(3, 7).value == 1
    assert ws.cell(4, 3).value == "1.2"
    assert ws.cell(4, 4).value == "Рама опорная"
