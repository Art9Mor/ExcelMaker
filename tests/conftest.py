"""
Модуль с фикстурами для тестов.
"""

import pytest
from pathlib import Path
from openpyxl import Workbook
from src.core.models import SpecDocument, Section, SpecItem, SpecHeader


@pytest.fixture
def sample_workbook_path(tmp_path: Path) -> Path:
    """
    Фикстура: создание временного Excel-файла с тестовыми данными.
    """

    wb = Workbook()
    ws = wb.active
    ws.title = "ВВ"

    # Заголовки
    ws.append([
        "Структура", "Скрыть строку, символы /*", "Наименование",
        "Цена, руб", "Количество", "Стоиомсть, руб", "Стоимость с наценкой, руб", "Ед. изм."
    ])

    # Секция Корпус
    ws.append(["Корпус", "", "", "", "", "", "", ""])
    ws.append(["", "", "Каркас ЯКНО", "80000", "1", "", "", "шт"])
    ws.append(["", "", "Рама опорная", "5000", "1", "", "", "шт"])

    # Секция Отсек
    ws.append(["Отсек высоковольтного выключателя", "", "", "", "", "", "", ""])
    ws.append(["", "", "Вакуумный выключатель", "145000", "1", "", "", "шт"])

    # Итог
    ws.append(["ИТОГ", "", "", "", "", "", "", ""])

    path = tmp_path / "sample.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def sample_spec_document():
    """
    Фикстура: создание тестового документа спецификации.
    """

    doc = SpecDocument(source_sheet="ВВ", target_sheet="Спецификация")
    doc.header = SpecHeader(equipment_type="ЯКНО-ВВ-6кВ")

    section = Section(number=1, title="Корпус")
    section.items.append(SpecItem(number="1.1", name="Каркас ЯКНО", unit="шт", quantity=1, price=80000, total=80000))
    section.items.append(SpecItem(number="1.2", name="Рама опорная", unit="шт", quantity=1, price=5000, total=5000))
    section.section_total = 85000
    doc.sections.append(section)
    doc.grand_total = 85000

    return doc


@pytest.fixture
def empty_workbook():
    """
    Фикстура: пустой workbook.
    """

    return Workbook()


@pytest.fixture
def workbook_with_headers():
    """
    Фикстура: workbook с заголовками.
    """

    wb = Workbook()
    ws = wb.active
    ws.append(["Структура", "", "Наименование", "Цена, руб", "Количество", "", "", "Ед. изм."])
    return wb
